from __future__ import annotations

from io import BytesIO
import re
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from app.services.roster_import import RosterImportError, parse_roster_workbook


def legacy_roster_export():
    """Anonymous workbook generated in memory; no local/student-data fixture needed."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["课程名称:计算机组成原理,课程号:211530001"])
    sheet.append(["学年：2025-2026 学期：2"])
    sheet.append(["序号", "学号", "姓名", "班级"])
    for index in range(1, 22):
        sheet.append([index, f"323195000{83 + index}", f"student_name_{index}", "物联（创）231"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream


def test_parse_students_xlsx_fixture():
    roster = parse_roster_workbook(legacy_roster_export())

    assert roster.course_name == "计算机组成原理"
    assert roster.course_code == "211530001"
    assert roster.semester == "2025-2026-2"
    assert len(roster.students) == 21
    assert roster.students[0].student_no == "01"
    assert roster.students[0].student_number == "32319500084"
    assert roster.students[0].student_name == "student_name_1"
    assert roster.students[0].source_class == "物联（创）231"


def exported_roster(*, wrong_dimension=False, cover=False, headers=None):
    workbook = Workbook()
    if cover:
        workbook.active.append(["封面"])
        sheet = workbook.create_sheet("名单")
    else:
        sheet = workbook.active
    sheet.append(["学年：2026-2027 学期：1"])
    sheet.append(["课程名称:操作系统,课程号:181520025"])
    sheet.append(headers or ["序号", "学号", "姓名", "性别", "专业", "班级"])
    sheet.append([1, "00001234", "测试学生", "", "", "测试班"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    if wrong_dimension:
        output = BytesIO()
        with ZipFile(stream) as source, ZipFile(output, "w") as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename.startswith("xl/worksheets/"):
                    data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', data)
                target.writestr(info, data)
        stream = output
    stream.seek(0)
    return stream


@pytest.mark.parametrize("wrong_dimension,cover", [(True, False), (False, True), (True, True)])
def test_teaching_export_ignores_stale_dimensions_and_finds_roster_sheet(wrong_dimension, cover):
    result = parse_roster_workbook(exported_roster(wrong_dimension=wrong_dimension, cover=cover))
    assert len(result.students) == 1
    assert result.students[0].student_no == "01"
    assert result.students[0].student_number == "00001234"
    assert result.course_name == "操作系统"
    assert result.course_code == "181520025"
    assert result.semester == "2026-2027-1"


def test_roster_header_aliases_and_whitespace():
    result = parse_roster_workbook(exported_roster(headers=["Ｎｏ．", "学 号", "学生姓名", "性别", "专业", "行政班级"]))
    assert result.students[0].source_class == "测试班"


def test_roster_invalid_file_and_missing_headers():
    with pytest.raises(RosterImportError, match="有效"):
        parse_roster_workbook(BytesIO(b"invalid workbook"))
    with pytest.raises(RosterImportError, match="序号"):
        parse_roster_workbook(exported_roster(headers=["备注"]))
